"""Build the reproducibility-oriented supplementary material DOCX."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "CXT-Fish_IMTS_Supplementary_Material_v2.docx"


def set_cell(cell, text: str, bold: bool = False, size: int = 8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, size=7):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True, size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): set_cell(cells[i], v, False, size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    return t


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.65); sec.right_margin = Inches(0.65)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(9)
    styles["Normal"].paragraph_format.space_after = Pt(2)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "Arial"
        styles[name].paragraph_format.space_before = Pt(6)
        styles[name].paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Supplementary Material\nCXT-Fish: Group-Disjoint Evaluation and Foreground-Sufficiency Training for Context-Robust Underwater Fish Recognition")
    r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("All supplementary analyses are based on frozen predictions, manifests, and metadata. No model was trained or re-inferred for this package; official Fish4Knowledge TEST was not accessed.").italic = True

    doc.add_heading("S1. Group-weighted robustness sensitivity", level=1)
    doc.add_paragraph("The manuscript's primary cross-class composite macro-F1 is class-balanced but image-weighted within species. This sensitivity gives each recorded group equal weight within species and then gives species equal weight. The nine fold–seed cells are averaged with equal weight. It is descriptive and does not replace the primary macro-F1 estimand.")
    s = pd.read_csv(ROOT / "experiments" / "cxt_fish_group_weighted_robustness_summary.csv")
    rows = []
    for _, x in s.iterrows():
        rows.append([x.metric, f"{x.F0_mean:.4f} ± {x.F0_sd:.4f}", f"{x.F1_mean:.4f} ± {x.F1_sd:.4f}", f"{x.delta_pp:+.2f}", f"{int(x.favourable_cells)}/{int(x.n_cells)}"])
    add_table(doc, ["Metric", "F0 mean ± SD", "CXT-Fish mean ± SD", "Δ (pp)", "Favourable cells"], rows, [2.2, 1.3, 1.6, 0.8, 1.0], size=8)
    doc.add_paragraph("Table S1. Cross-class donor-context-composite group-weighted sensitivity. SD is the sample standard deviation across nine equally weighted fold–seed cells. Official-test access: false.")

    doc.add_heading("S2. Recorded-group counts by species and outer fold", level=1)
    counts = pd.read_csv(ROOT / "experiments" / "cxt_fish_outer_recorded_group_counts_16x3.csv", dtype={"species_id": str})
    wide = counts.pivot(index="species_id", columns="fold", values="n_recorded_groups").reset_index()
    wide.columns = ["Species"] + [f"Fold {c}" for c in wide.columns[1:]]
    add_table(doc, list(wide.columns), wide.values.tolist(), [1.3, 1.1, 1.1, 1.1], size=8)
    doc.add_paragraph(f"Table S2. Recorded-group counts in the three frozen outer-test manifests. There are {len(counts)} fold × species strata; the minimum is {counts.n_recorded_groups.min()} and the maximum is {counts.n_recorded_groups.max()}. These finite clusters motivate cautious interpretation of percentile cluster intervals for sparsely represented strata.")

    doc.add_heading("S3. Donor-species audit", level=1)
    d = pd.read_csv(ROOT / "experiments" / "cxt_fish_donor_species_effects.csv")
    drows = []
    for _, x in d.iterrows():
        drows.append([x["donor_species"], int(x["n_recipients"]), f"{x['fraction_of_all_cross_composites']:.3f}", int(x["unique_donor_groups"]), f"{x['effect']*100:+.2f} pp"])
    add_table(doc, ["Donor species", "Recipients", "Fraction", "Donor groups", "Correctness effect"], drows, [1.2, 1.1, 0.9, 1.1, 1.3], size=7)
    doc.add_paragraph("Table S3. Frozen seed-3407 donor-species audit. The donor-species-equal-weight correctness sensitivity is +3.71 percentage points; this is descriptive and not a replacement for the natural-manifest-weighted main result.")

    p_s4 = doc.add_heading("S4. MobileNetV3-Large fold-level sensitivity", level=1)
    p_s4.paragraph_format.page_break_before = True
    m = pd.read_csv(ROOT / "experiments" / "cxt_fish_mobilenet_per_fold_results.csv")
    mrows = []
    for fold in sorted(m.fold.unique()):
        a = m[(m.fold == fold) & (m.method == "MV0")].iloc[0]
        b = m[(m.fold == fold) & (m.method == "MV1")].iloc[0]
        mrows.append([int(fold), f"{(b.original_macro_f1-a.original_macro_f1)*100:+.2f}", f"{(b.foreground_macro_f1-a.foreground_macro_f1)*100:+.2f}", f"{(b.cross_swap_macro_f1-a.cross_swap_macro_f1)*100:+.2f}", f"{(b.dar_flip-a.dar_flip)*100:+.2f}"])
    add_table(doc, ["Fold", "Clean Δ pp", "Foreground Δ pp", "Cross composite Δ pp", "DAR-flip Δ pp"], mrows, [0.8, 1.1, 1.4, 1.6, 1.2], size=8)
    doc.add_paragraph("Table S4. MobileNetV3-Large architecture-sensitivity fold effects at the pre-specified seed 3407. Foreground improves in 3/3 folds, cross-class composite improves in 2/3 folds, and DAR-flip decreases in 3/3 folds; clean effects are heterogeneous and negative on average.")

    doc.add_heading("S5. Evidence-role and reproducibility notes", level=1)
    notes = [
        "Main ResNet18 evidence: same-corpus frozen group-disjoint re-evaluation, 3 folds × 3 seeds (17, 2026, 3407).",
        "Primary estimand: equal-weight mean of nine unrounded fold–seed cross-class macro-F1 differences.",
        "Primary interval: 5,000 paired species-stratified group-cluster bootstrap replicates, conditional on completed development, frozen models/seeds, and frozen seed-3407 donor realization.",
        "Donor-realization, F0-2RGB, donor-subject-suppressed, group-weighted, and MobileNet analyses: post-hoc or boundary evidence; not independent validation.",
        "Route C: fixed mechanism stress control, not faithful CLIB reproduction and not an exhaustive contrastive-learning comparison; detailed values are in Supplementary Note S6.",
        "Higher acquisition units (camera/session/deployment/date) were not verifiable from frozen recognition metadata.",
        "Official Fish4Knowledge TEST accessed: false.",
        "The repository retains stopped TCCR, DTH, donor-aware, TRAFS, CIR, DLE, TAP, and CXT-Select explorations as negative-result records.",
    ]
    for n in notes: doc.add_paragraph(n, style="List Bullet")
    doc.add_heading("S6. Route C mechanism stress control", level=1)
    doc.add_paragraph("Supplementary Note S6. Route C is a separately frozen two-stage mask-guided subject–non-primary contrastive mechanism control. It completed 9/9 fold–seed cells with official-test access false in every metrics file. Its equal-weight means were 0.6984 original-view macro-F1, 0.6981 foreground macro-F1, 0.7014 same-class-composite macro-F1, 0.4224 cross-class-composite macro-F1, and 0.1610 DAR-flip. Route C is retained as a fixed stress control and is not interpreted as a faithful CLIB reproduction or an exhaustive contrastive-learning comparison.")

    doc.add_heading("S7. Donor residual-bin heterogeneity", level=1)
    residual = pd.read_csv(ROOT / "experiments" / "cxt_fish_effect_vs_donor_residual.csv")
    rrows = []
    for _, x in residual[residual.n_cell_observations > 0].iterrows():
        rrows.append([x["bin"], int(x["n_cell_observations"]), int(x["n_unique_recipients"]), f"{x['f0_correctness']:.3f}", f"{x['cxt_correctness']:.3f}", f"{x['paired_correctness_delta']*100:+.2f} pp", f"{x['dar_flip_f0']:.3f}", f"{x['dar_flip_cxt']:.3f}"])
    add_table(doc, ["Residual bin", "Pairs", "Recipients", "F0 correctness", "CXT-Fish correctness", "Delta (pp)", "F0 DAR-flip", "CXT DAR-flip"], rrows, [1.1, 0.8, 0.9, 1.1, 1.3, 0.9, 0.9, 1.0], size=7)
    doc.add_paragraph("Table S7. Frozen donor-subject residual-bin sensitivity. Bins are defined by donor foreground fraction outside the recipient hard mask; empty bins are not shown. Values are descriptive and do not define a new inferential endpoint.")

    doc.add_heading("S8. Donor-species equal-weight sensitivity", level=1)
    s8rows = []
    for _, x in d.iterrows():
        s8rows.append([x["donor_species"], int(x["n_recipients"]), int(x["unique_donor_groups"]), f"{x['f0_correct']:.3f}", f"{x['cxt_correct']:.3f}", f"{x['effect']*100:+.2f} pp"])
    add_table(doc, ["Donor species", "Recipients", "Donor groups", "F0 correctness", "CXT-Fish correctness", "Delta (pp)"], s8rows, [1.1, 1.0, 1.0, 1.1, 1.3, 0.9], size=6)
    doc.add_paragraph("Table S8. Donor-species equal-weight correctness sensitivity. Species are weighted equally for this descriptive audit; the natural-manifest-weighted main result remains primary.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
